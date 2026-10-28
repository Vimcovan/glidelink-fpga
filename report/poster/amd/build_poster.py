"""AMD poster template -> GlideLink poster (one 13.33 x 7.5 in slide).

Keeps the template frame (title, team, AMD logo, product photo, INTRODUCTION header, CREATIVE DESIGN and RESULT
bars, the dividing lines) and lays out the content freely in the three columns.

    python build_poster.py AMD_poster_template.pptx ../poster_amd.pptx
    python render.py ../poster_amd.pptx ../poster_amd_preview.png 2400 ../poster_amd.pdf   (PowerPoint COM)

team.json, uestc_logo.png and strip.png (decoded-image rows of data/measurements/psnr_snr/k0802_final_visual.png)
live next to this script.
"""
import copy, json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.normpath(os.path.join(HERE, '..', '..', '..'))
TEMPLATE, OUT = sys.argv[1], sys.argv[2]
INFO = json.load(open(os.path.join(HERE, 'team.json'), encoding='utf-8'))
BLUE, ORANGE = RGBColor(0x2E, 0x5E, 0x8E), RGBColor(0xC8, 0x60, 0x2A)
INK, MUTE, BLACK = RGBColor(0x1F, 0x29, 0x33), RGBColor(0x52, 0x60, 0x6D), RGBColor(0, 0, 0)
PS_F, PS_L = RGBColor(0xEE, 0xE8, 0xF5), RGBColor(0x7E, 0x6A, 0x9E)
PL_F, PL_L = RGBColor(0xDD, 0xEB, 0xF7), BLUE
RF_F, RF_L = RGBColor(0xE8, 0xF2, 0xE1), RGBColor(0x5B, 0x8C, 0x3A)
BODY = 11

prs = Presentation(TEMPLATE)
s = prs.slides[0]
by_id = {sh.shape_id: sh for sh in s.shapes}


def remove(sh):
    sh._element.getparent().remove(sh._element)


def set_paras(sh, paras):
    """Replace the text of a template shape, keeping its first paragraph and run style."""
    txb = sh._element.find(qn('p:txBody'))
    ps = txb.findall(qn('a:p'))
    ppr = ps[0].find(qn('a:pPr'))
    rpr = next(p.find(qn('a:r')) for p in ps if p.find(qn('a:r')) is not None).find(qn('a:rPr'))
    for h in rpr.findall(qn('a:highlight')):
        rpr.remove(h)
    for p in ps:
        txb.remove(p)
    for text in paras:
        p = etree.SubElement(txb, qn('a:p'))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn('a:r'))
        r.append(copy.deepcopy(rpr))
        etree.SubElement(r, qn('a:t')).text = text
    return sh


def box(x, y, w, h, paras, size=BODY, align=PP_ALIGN.LEFT, bold=False, color=BLACK, italic=False,
        space=0, indent=False, anchor=MSO_ANCHOR.TOP):
    """paras: list of paragraphs; a paragraph is a str or a list of runs (str or (text, opts))."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if not isinstance(paras, list):
        paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if k:
            p.space_before = Pt(space)
        if indent:
            pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(Inches(0.18))); pPr.set('indent', str(-Inches(0.18)))
        for run in (para if isinstance(para, list) else [para]):
            text, o = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get('sz', size)); f.bold = o.get('b', bold); f.italic = o.get('i', italic)
            f.name = 'Calibri'; f.color.rgb = o.get('color', color)
    return tb


def block(x, y, w, h, title, sub, fill, line, tsz=9, ssz=7.5):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.12
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for k, (text, sz, b, col) in enumerate(((title, tsz, True, INK), (sub, ssz, False, MUTE))):
        if not text:
            continue
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text
        r.font.size = Pt(sz); r.font.bold = b; r.font.name = 'Calibri'; r.font.color.rgb = col
    return sh


def arrow(x1, y1, x2, y2, color=INK, dash=False, width=1.0, both=False, head='sm'):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn('a:prstDash')).set('val', 'dash')
    te = etree.SubElement(ln, qn('a:tailEnd'))
    te.set('type', 'triangle'); te.set('w', head); te.set('len', head)
    if both:
        he = etree.Element(qn('a:headEnd')); he.set('type', 'triangle'); he.set('w', head); he.set('len', head)
        te.addprevious(he)
    return c


def picture(path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


# ---------------------------------------------------------------- template frame
for name in ('TextBox 5', 'TextBox 48', 'TextBox 50'):          # master: "[Copyright <Year>]", slide number
    for sh in list(prs.slide_master.shapes):
        if sh.name == name:
            remove(sh)
for i in list(range(48, 57)) + [7, 11, 40, 41, 42, 43, 44, 45, 46, 47]:
    remove(by_id[i])                                             # hints and the sample content

t = set_paras(by_id[9], ['GlideLink: Real-Time DeepJSCC-Q', 'Image Transmission on AMD FPGA'])
t.left, t.top, t.width = Inches(0.23), Inches(0.16), Inches(6.3)
for p in t.text_frame.paragraphs:
    p.line_spacing = 0.9
by_id[8].top = Inches(1.1)                                       # INTRODUCTION
set_paras(by_id[19], ['G']); set_paras(by_id[20], ['N'])              # the template's vertical label reads DESING
nm = set_paras(by_id[29], [INFO['members']])
nm.left, nm.width = Inches(6.95), Inches(3.3)
af = set_paras(by_id[31], INFO['affiliation'])
af.left, af.width = Inches(7.72), Inches(2.5)
if INFO.get('badge'):
    picture(os.path.join(HERE, INFO['badge']), 10.22, 0.14, h=0.82)
picture(os.path.join(REPO, 'report', 'poster', 'assets', 'hardware_cutout.jpg'), 11.06, 0.1, w=2.24)   # amd/photo_cutout.py
cap = set_paras(by_id[30], ['On-board test: RX (left), TX (right)', 'AMD Zynq UltraScale+ MPSoC'])
cap.left, cap.top, cap.width, cap.height = Inches(11.0), Inches(1.36), Inches(2.3), Inches(0.4)
for p in cap.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(0)

# ---------------------------------------------------------------- INTRODUCTION (x 0.2 .. 4.25)
LX, LW = 0.22, 4.0
box(LX, 1.68, LW, 1.3, [[
    ('Live video from drones, robots and remote cameras must cross radio links whose SNR changes from moment to '
     'moment. ', {}),
    ('Separate source-channel coding (SSCC)', {'b': True}),
    (' compresses the image with JPEG and adds a channel code: below the code threshold, bit errors propagate '
     'through the variable-length code and the whole image is lost (the ', {}),
    ('cliff effect', {'b': True}), ('); above it, quality no longer improves.', {})]])
box(LX, 3.0, LW, 0.22, [[('Typical application case', {'b': True})]], size=12, align=PP_ALIGN.CENTER)
CH = [('Camera', 'drone / robot'), ('Encoder', 'FPGA, TX'), ('Radio channel', 'noise, fading'),
      ('Decoder', 'FPGA, RX'), ('Display', 'live view')]
CF = [(PS_F, PS_L), (PL_F, PL_L), (RF_F, RF_L), (PL_F, PL_L), (PS_F, PS_L)]
x0, w, g, y = 0.3, 0.64, 0.17, 3.27
for k, (a, b) in enumerate(CH):
    block(x0 + k * (w + g), y, w, 0.56, a, b, *CF[k], tsz=8.5, ssz=7)
    if k < 4:
        arrow(x0 + k * (w + g) + w, y + 0.28, x0 + (k + 1) * (w + g), y + 0.28)
box(LX, 3.98, LW, 3.15, [
    [('Deep joint source-channel coding (DeepJSCC)', {'b': True}),
     (' maps pixels directly to channel symbols with a pair of CNNs, so image quality follows the SNR smoothly. ', {}),
     ('DeepJSCC-Q', {'b': True}),
     (' (Tung et al., JSAIT 2022) restricts the symbols to a QAM constellation that a standard digital PHY can '
      'carry, but it is computation-heavy and the receiver must pass undecided soft symbols to the decoder.', {})],
    [('GlideLink', {'b': True}),
     (' builds a complete real-time link on two AMD Zynq UltraScale+ MPSoC boards (PYNQ-ZU) with AD9361 radios. '
      'The DeepJSCC-Q encoder, the decoder and an 802.11a-like OFDM PHY all run in the PL and send ', {}),
     ('256 × 256 RGB images at 30 fps over the air', {'b': True}),
     ('. The PS captures camera frames, sets up the radios and shows the live link on a touch screen.', {})],
    [('Key technical issues: ', {'b': True}),
     ('fitting a 1.15-GMAC-per-frame CNN codec into one FPGA at 30 fps with all weights on chip; synchronizing and '
      'tracking 683-symbol OFDM frames precisely enough for soft-symbol decoding; and a fair over-the-air comparison '
      'with JPEG + LDPC on the same PHY.', {})],
], space=7)

# ---------------------------------------------------------------- CREATIVE DESIGN (x 4.35 .. 8.3)
MX = 4.38
# system composition: two boards (PS + PL), RF cards, peripherals and the links between them
GREY = RGBColor(0x8A, 0x96, 0xA3)
FR_F = RGBColor(0xF6, 0xF8, 0xFA)


def frame(x, y, w, h, title):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = FR_F
    sh.line.color.rgb = GREY; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    box(x + 0.06, y + 0.03, w - 0.12, 0.16, [title], size=7.5, align=PP_ALIGN.RIGHT)


def label(x, y, w, text, align=PP_ALIGN.CENTER, color=MUTE, size=6.5):
    box(x, y, w, 0.13, [[(text, {'color': color})]], size=size, align=align)


def antenna(x, y):                                   # mast from (x, y) upwards with a V on top
    for x2, y2 in ((x, y - 0.22), (x - 0.06, y - 0.3), (x + 0.06, y - 0.3)):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y - (0 if x2 == x else 0.22)),
                                   Inches(x2), Inches(y2))
        c.line.color.rgb = INK; c.line.width = Pt(1)


PX, PW = MX, 0.56                                    # peripherals
FX, FW = PX + PW + 0.2, 1.9                          # board frame
SX, IW = FX + 0.05, 0.78                             # PS and PL boxes inside the frame
LX2 = SX + IW + 0.24
RX_, RW = FX + FW + 0.2, 0.56                        # RF card
for yf, tx in ((1.42, True), (2.62, False)):
    yb, hb = yf + 0.2, 0.58
    frame(FX, yf, FW, 0.84, [('TX board' if tx else 'RX board', {'b': True}), ('  PYNQ-ZU (xczu5eg)', {'color': MUTE})])
    if tx:
        block(SX, yb, IW, hb, 'PS', '4× Cortex-A53, Linux: camera, radio setup', PS_F, PS_L, tsz=8, ssz=6.5)
        block(LX2, yb, IW, hb, 'PL', 'CNN encoder, OFDM TX, LDPC encoder', PL_F, PL_L, tsz=8, ssz=6.5)
        block(PX, yb, PW, hb, 'USB camera', None, RGBColor(0xFF, 0xFF, 0xFF), GREY, tsz=8)
        arrow(PX + PW, yb + hb / 2, SX, yb + hb / 2); label(PX + PW - 0.1, yb + hb / 2 - 0.15, 0.4, 'USB')
    else:
        block(SX, yb, IW, hb, 'PS', 'A53 + Mali-400 GPU: GUI, web page', PS_F, PS_L, tsz=8, ssz=6.5)
        block(LX2, yb, IW, hb, 'PL', 'OFDM RX, CNN decoder, LDPC decoder', PL_F, PL_L, tsz=8, ssz=6.5)
        block(PX, yb, PW, 0.27, 'Touch screen', None, RGBColor(0xFF, 0xFF, 0xFF), GREY, tsz=7.5)
        block(PX, yb + 0.31, PW, 0.27, 'Phone', None, RGBColor(0xFF, 0xFF, 0xFF), GREY, tsz=7.5)
        arrow(SX, yb + 0.135, PX + PW, yb + 0.135); label(PX + PW - 0.1, yb + 0.135 - 0.15, 0.4, 'DP')
        arrow(SX, yb + 0.445, PX + PW, yb + 0.445, dash=True, both=True); label(PX + PW - 0.1, yb + 0.445 - 0.15, 0.4, 'Wi-Fi')
    block(RX_, yb, RW, hb, 'AD9361', 'RF card', RF_F, RF_L, tsz=8, ssz=6)
    if tx:
        arrow(SX + IW, yb + hb / 2, LX2, yb + hb / 2, both=True)
        arrow(LX2 + IW, yb + hb / 2, RX_, yb + hb / 2)
    else:
        arrow(LX2, yb + hb / 2, SX + IW, yb + hb / 2, both=True)
        arrow(RX_, yb + hb / 2, LX2 + IW, yb + hb / 2)
    label(SX + IW - 0.06, yb + hb / 2 - 0.15, 0.36, 'AXI')
    label(LX2 + IW - 0.06, yb + hb / 2 - 0.15, 0.4, 'FMC')
    antenna(RX_ + RW + 0.07, yb + 0.32)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(RX_ + RW), Inches(yb + 0.32), Inches(RX_ + RW + 0.07), Inches(yb + 0.32))
    c.line.color.rgb = INK; c.line.width = Pt(1)
arrow(RX_ + RW / 2, 1.42 + 0.2 + 0.58, RX_ + RW / 2, 2.62 + 0.2, color=ORANGE, width=1.75)
label(RX_ + RW / 2 + 0.04, 2.3, 0.8, 'over the air', align=PP_ALIGN.LEFT, color=ORANGE, size=7)
arrow(SX + IW / 2, 1.42 + 0.2 + 0.58, SX + IW / 2, 2.62 + 0.2, color=MUTE, dash=True, both=True)
label(SX + IW / 2 + 0.05, 2.3, 1.4, 'Ethernet: TX preview, control', align=PP_ALIGN.LEFT, size=7)
sa = box(7.95, 1.45, 0.3, 2.0, [[('System Architecture', {'b': True})]], size=12, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
sa.text_frame._txBody.find(qn('a:bodyPr')).set('vert', 'vert')

box(MX, 3.57, 3.95, 0.25, [[('Key innovations and technical issues addressed:', {'b': True})]], size=13)
gui = picture(os.path.join(REPO, r'report\poster\assets\gui_phy_page.png'), 5.96, 5.9, w=2.32)
box(MX, 3.9, 3.95, 2.2, [
    [('1. Streaming CNN accelerator: ', {'b': True}),
     ('W8A12 integer layers on 64 convolution engines in a layer pipeline at 250 MHz; all weights and '
      'activations stay on chip, no DDR traffic.', {})],
    [('2. Asymmetric shallow decoder: ', {'b': True}),
     ('trained jointly with the encoder, it needs 1/102 of the original multiply-accumulates. The codec does 5.2× '
      'the workload of a published FPGA DeepJSCC (GLOBECOM 2025) with fewer DSPs and BRAMs.', {})],
    [('3. Soft-symbol OFDM PHY: ', {'b': True}),
     ('683-symbol frames with PN scrambling and second-order SFO/CPE tracking; equalized I/Q go to the '
      'decoder without hard decisions.', {})],
], space=5)
box(MX, 5.98, 1.5, 1.2, [[('4. Live system: ', {'b': True}),
     ('camera to touch screen with live PSNR, constellations and a 3-D channel view on the Mali-400 GPU.', {})]])

# ---------------------------------------------------------------- RESULT (x 9.2 .. 13.25)
da = json.load(open(os.path.join(REPO, r'data\measurements\psnr_snr\k0802_final_da.json'), encoding='utf-8'))
cd = XyChartData()
for mode, name in (('jscc', 'DeepJSCC-Q'), ('sscc', 'JPEG + LDPC (SSCC)')):
    pts = sorted((r['snr_da_db'], r['psnr_shown']) for r in da['results']
                 if r['mode'] == mode and r.get('snr_da_db') is not None and r.get('psnr_shown') is not None
                 and (r.get('snr_da_excl') or 0) < 0.05 and not (mode == 'jscc' and (r.get('rx_fps') or 0) < 25))
    ser = cd.add_series(name)
    for xv, yv in pts:
        ser.add_data_point(round(xv, 2), round(yv, 2))
box(9.8, 2.12, 3.45, 0.2, [[('Over-the-air PSNR (DIV2K 0802, 32,768 64QAM symbols per image)', {'b': True})]], size=8)
gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(9.75), Inches(2.3), Inches(3.5), Inches(1.8), cd)
ch = gf.chart
ch.font.size = Pt(8); ch.font.name = 'Calibri'
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
for ser, col, mk in zip(ch.plots[0].series, (BLUE, ORANGE), (XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.TRIANGLE)):
    ser.format.line.color.rgb = col; ser.format.line.width = Pt(1.5); ser.smooth = False
    ser.marker.style = mk; ser.marker.size = 4
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    ser.marker.format.line.color.rgb = col
xa, ya = ch.category_axis, ch.value_axis
xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, 39, 6
ya.minimum_scale, ya.maximum_scale, ya.major_unit = 5, 35, 5
for ax, title in ((xa, 'Over-the-air SNR (dB)'), (ya, 'PSNR (dB)')):
    ax.has_title = True
    ax.axis_title.text_frame.text = title
    r = ax.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(8); r.font.bold = False; r.font.color.rgb = INK
    ax.tick_labels.font.size = Pt(8); ax.tick_labels.font.color.rgb = MUTE
    ax.format.line.color.rgb = RGBColor(0xB8, 0xC7, 0xD6)
ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE8, 0xEE)
xa.has_major_gridlines = False

box(9.25, 4.17, 4.0, 0.62, [[
    ('DeepJSCC-Q is ', {}), ('2.8 dB', {'b': True}), (' above JPEG + LDPC at high SNR and still gives ', {}),
    ('23.5 dB at 2.9 dB SNR', {'b': True}), (', while the baseline fails abruptly at 16 to 19.5 dB.', {})]])
picture(os.path.join(HERE, 'strip.png'), 9.25, 4.78, w=4.0)
box(9.25, 6.38, 4.0, 0.2, [[('30 fps', {'b': True}), (' for 10 min (18,020 frames),  ', {}),
                             ('67.5 ms', {'b': True}), (' PL end-to-end latency', {})]], size=10)
ROWS = [('Design', 'LUT', 'BRAM36', 'DSP', 'WNS (ns)', 'Power*'),
        ('TX', '54,241', '104.5', '243', '+0.082', '3.07 W'),
        ('RX', '63,699', '65.5', '133', '+0.307', '3.44 W')]
tbl = s.shapes.add_table(3, 6, Inches(9.25), Inches(6.6), Inches(4.0), Inches(0.42)).table
for c, wd in enumerate((0.55, 0.75, 0.7, 0.5, 0.75, 0.75)):
    tbl.columns[c].width = Inches(wd)
for r_, row in enumerate(ROWS):
    tbl.rows[r_].height = Inches(0.14)
    for c, val in enumerate(row):
        cell = tbl.cell(r_, c)
        cell.margin_left = cell.margin_right = Inches(0.03); cell.margin_top = cell.margin_bottom = 0
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xDD, 0xEB, 0xF7) if r_ == 0 else RGBColor(0xFF, 0xFF, 0xFF)
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = val
        run.font.size = Pt(8.5); run.font.name = 'Calibri'; run.font.bold = (r_ == 0); run.font.color.rgb = BLACK
tbl.first_row = True
box(9.25, 7.04, 3.1, 0.15, '* measured, PS + PL + DDR; xczu5eg after routing', size=7, color=MUTE)
prs.save(OUT)
print('saved', OUT)
